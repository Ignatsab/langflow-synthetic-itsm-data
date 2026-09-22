from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
GUIDES = {
    "data/automatic_rule_provisioning/README.md": "data/automatic_rule_provisioning/Automatic_Rule_Provisioning_Data_Guide.docx",
    "data/pam_chatbot/README.md": "data/pam_chatbot/PAM_Chatbot_Data_Guide.docx",
    "data/automated_vulnerability_management/README.md": "data/automated_vulnerability_management/Automated_Vulnerability_Management_Data_Guide.docx",
    "data/aps_ai_knowledge_base/README.md": "data/aps_ai_knowledge_base/APS_AI_Knowledge_Base_Data_Guide.docx",
    "data/autosys_sod_eta/README.md": "data/autosys_sod_eta/AutoSys_SOD_ETA_Data_Guide.docx",
}

NAVY = "17365D"
PALE_BLUE = "EAF1F8"
LIGHT_GRAY = "D9D9D9"
TEXT_GRAY = RGBColor(89, 89, 89)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), LIGHT_GRAY)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_table_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def sanitize_heading(text: str) -> str:
    text = text.replace("`", "")
    text = re.sub(r"\.json\b", " JSON", text, flags=re.I)
    text = re.sub(r"\.py\b", " Python Generator", text, flags=re.I)
    text = text.replace("0-100", "0 to 100")
    text = text.replace("_", " ")
    text = re.sub(r"[-–—/:;,.()]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    replacements = {
        "aps": "APS", "ai": "AI", "pam": "PAM", "cve": "CVE", "json": "JSON",
        "autosys": "AutoSys", "sod": "SOD", "eta": "ETA", "reqiwav": "REQIWAV",
        "servicenow": "ServiceNow", "rag": "RAG", "cmd": "CMD", "box": "BOX",
    }
    words = [replacements.get(word.lower(), word) for word in text.split()]
    result = " ".join(words)
    return result[:1].upper() + result[1:]


def add_inline(paragraph, text: str, base_bold: bool = False) -> None:
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(45, 45, 45)
        else:
            run = paragraph.add_run(token)
            run.bold = base_bold


def configure_document(document: Document, title: str, source_path: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(5)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 16, 15, 7),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 11.5, 10, 4),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor(32, 32, 32)
    code.paragraph_format.left_indent = Inches(0.25)
    code.paragraph_format.right_indent = Inches(0.15)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    header.text = "Hackathon Participant Data Guide"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8.5)
        run.font.color.rgb = TEXT_GRAY

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Page ")
    run.font.name = "Aptos"
    run.font.size = Pt(8.5)
    run.font.color.rgb = TEXT_GRAY
    add_field(footer, "PAGE")
    run = footer.add_run(" of ")
    run.font.name = "Aptos"
    run.font.size = Pt(8.5)
    run.font.color.rgb = TEXT_GRAY
    add_field(footer, "NUMPAGES")

    title_paragraph = document.add_paragraph(style="Title")
    title_ppr = title_paragraph._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    title_paragraph.add_run(sanitize_heading(title))
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(3)
    run = subtitle.add_run("Hackathon Participant Data Guide")
    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(11)
    run.font.color.rgb = NAVY and RGBColor(23, 54, 93)
    source = document.add_paragraph()
    source.paragraph_format.space_after = Pt(12)
    run = source.add_run(f"Source README  {source_path}")
    run.font.name = "Aptos"
    run.font.size = Pt(8.5)
    run.font.color.rgb = TEXT_GRAY


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    available = 7.06
    if columns == 2:
        widths = [2.0, 5.06]
    elif columns == 3:
        widths = [1.65, 1.05, 4.36]
    else:
        widths = [available / columns] * columns
    for row_index, values in enumerate(rows):
        for col_index in range(columns):
            cell = table.cell(row_index, col_index)
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            text = values[col_index].strip() if col_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, text, base_bold=row_index == 0)
            for run in paragraph.runs:
                run.font.name = "Aptos" if run.font.name != "Consolas" else "Consolas"
                run.font.size = Pt(9 if row_index else 9.2)
                if row_index == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
        if row_index == 0:
            repeat_table_header(table.rows[row_index])
        keep_table_row_together(table.rows[row_index])
    set_table_borders(table)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if index == start + 1 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            index += 1
            continue
        rows.append(cells)
        index += 1
    return rows, index


def build_guide(readme_path: Path, output_path: Path) -> None:
    lines = readme_path.read_text(encoding="utf-8").splitlines()
    source_title = lines[0].lstrip("# ").strip()
    document = Document()
    configure_document(document, source_title, str(readme_path.relative_to(ROOT)))

    index = 1
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph(style="Code Block")
                paragraph.add_run("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = min(max(len(heading.group(1)) - 1, 1), 3)
            document.add_paragraph(sanitize_heading(heading.group(2)), style=f"Heading {level}")
            index += 1
            continue
        if stripped.endswith(":") and len(stripped) <= 60:
            document.add_paragraph(sanitize_heading(stripped), style="Heading 3")
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[index + 1]):
            rows, index = parse_table(lines, index)
            add_markdown_table(document, rows)
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            next_is_bullet = index + 1 < len(lines) and bool(re.match(r"^\s*[-*]\s+", lines[index + 1]))
            following_is_bullet = index + 2 < len(lines) and bool(re.match(r"^\s*[-*]\s+", lines[index + 2]))
            if next_is_bullet and not following_is_bullet:
                paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, bullet.group(1))
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, numbered.group(1))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "```", "|", "- ", "* ")) or re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))

    document.core_properties.title = sanitize_heading(source_title)
    document.core_properties.subject = "Hackathon participant dataset field guide"
    document.core_properties.author = "Hackathon Data Team"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    for source, target in GUIDES.items():
        readme_path = ROOT / source
        output_path = ROOT / target
        build_guide(readme_path, output_path)
        print(output_path.relative_to(ROOT))


if __name__ == "__main__":
    main()
